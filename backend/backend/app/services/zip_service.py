"""
ZIP Archive Ingestion Service
-----------------------------
Handles extraction of nested ZIP archives containing audit reports:
- Traverses nested folder hierarchies safely (zip-slip prevention)
- Harvests candidate passwords from:
    1. Text files (.txt, .csv, .tsv)
    2. Word files (.docx) containing passwords, credentials, or mapping tables
    3. Directory-scoped password files
- Maps passwords hierarchically:
    1. Filename-specific mappings (e.g. 'audit1.pdf: Secret123')
    2. Directory-scoped passwords (passwords found in the same folder as the report)
    3. Ancestor / root directory passwords
    4. Global archive passwords
- Pre-emptively attempts decryption of password-protected reports (.pdf, .docx)
- Filters out non-report artifacts (images, credential files, system files)
"""

import io
import os
import re
import zipfile
import uuid
import logging
from typing import List, Dict, Tuple, Optional, Set

import docx

from app.utils.encryption_utils import (
    is_pdf_encrypted,
    is_docx_encrypted,
    decrypt_pdf,
    decrypt_docx,
)
from app.utils.file_utils import validate_extension

logger = logging.getLogger(__name__)

VALID_REPORT_EXTENSIONS = {".pdf", ".docx", ".doc"}
VALID_TEXT_EXTENSIONS = {".txt", ".csv", ".tsv"}

# Keywords indicating that a docx/txt file contains passwords/credentials
PWD_FILE_KEYWORDS = {
    "password", "passwords", "pwd", "pwds", "credential", "credentials",
    "secret", "secrets", "passcode", "passcodes", "pass_word", "pass-word",
    "key", "keys", "pin", "pins"
}


class ZipExtractedReport:
    """Represents a report file extracted from a ZIP archive."""
    def __init__(
        self,
        relative_path: str,
        filename: str,
        file_bytes: bytes,
        protection_status: str = "No",  # "No", "Decrypted", "Locked"
    ):
        self.relative_path = relative_path
        self.filename = filename
        self.file_bytes = file_bytes
        self.protection_status = protection_status


def extract_passwords_from_text(content: str) -> Tuple[List[str], Dict[str, str]]:
    """
    Parses plain text content to extract passwords and key-value file mappings.
    Returns:
        (candidate_passwords_list, filename_to_password_dict)
    """
    candidates: List[str] = []
    name_map: Dict[str, str] = {}
    seen_candidates: Set[str] = set()

    def add_candidate(pwd: str):
        pwd = pwd.strip().strip('"').strip("'")
        if not pwd or len(pwd) > 128:
            return
        low = pwd.lower()
        # Filter out common header words
        if low in {
            "password", "passwords", "pwd", "pwds", "pass", "key", "keys",
            "pin", "filename", "file", "report", "name", "id", "status"
        }:
            return
        if pwd not in seen_candidates:
            seen_candidates.add(pwd)
            candidates.append(pwd)

    for raw_line in content.splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or line.startswith("//"):
            continue

        delim_matched = False
        # Check delimiters for key-value pairs (colon, equals, comma, tab)
        for delim in (":", "=", "\t", ","):
            if delim in line:
                parts = line.split(delim, 1)
                left = parts[0].strip().strip('"').strip("'")
                right = parts[1].strip().strip('"').strip("'")
                left_lower = left.lower()

                # If left is generic indicator like 'password', 'pwd'
                if any(kw in left_lower for kw in ("password", "pwd", "passcode", "secret", "pin", "key")):
                    if right:
                        add_candidate(right)
                        delim_matched = True
                        break

                # If right looks like a plausible password
                if 1 <= len(right) <= 128:
                    base_name = os.path.splitext(left)[0].strip().lower()
                    if base_name and base_name not in {"file", "filename", "report", "document", "name"}:
                        name_map[base_name] = right
                        name_map[left.lower()] = right
                        add_candidate(right)
                        delim_matched = True
                        break

        if not delim_matched:
            # Single candidate password on this line
            add_candidate(line)

    return candidates, name_map


def extract_passwords_from_docx(file_bytes: bytes) -> Tuple[List[str], Dict[str, str]]:
    """
    Extracts passwords and file mappings from a Word (.docx) document.
    Scans tables (column headers & key-value rows) and text paragraphs.
    """
    candidates: List[str] = []
    name_map: Dict[str, str] = {}
    seen_candidates: Set[str] = set()

    def merge_results(c_list: List[str], m_dict: Dict[str, str]):
        for c in c_list:
            if c not in seen_candidates:
                seen_candidates.add(c)
                candidates.append(c)
        name_map.update(m_dict)

    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        logger.warning(f"Failed to parse docx for password extraction: {e}")
        return candidates, name_map

    # 1. Inspect tables
    for table in doc.tables:
        if not table.rows:
            continue

        header_cells = [cell.text.strip().lower() for cell in table.rows[0].cells]
        file_col_idx = -1
        pass_col_idx = -1

        for idx, col_name in enumerate(header_cells):
            if any(k in col_name for k in ("file", "report", "audit", "document", "name", "id")):
                if file_col_idx == -1:
                    file_col_idx = idx
            if any(k in col_name for k in ("pass", "pwd", "key", "code", "secret", "pin")):
                if pass_col_idx == -1:
                    pass_col_idx = idx

        start_row = 1 if (file_col_idx != -1 or pass_col_idx != -1) else 0

        for row in table.rows[start_row:]:
            cells = [c.text.strip() for c in row.cells]
            if not cells:
                continue

            if file_col_idx != -1 and pass_col_idx != -1 and file_col_idx < len(cells) and pass_col_idx < len(cells):
                f_val = cells[file_col_idx]
                p_val = cells[pass_col_idx]
                if f_val and p_val:
                    base = os.path.splitext(f_val)[0].strip().lower()
                    name_map[base] = p_val
                    name_map[f_val.lower()] = p_val
                    if p_val not in seen_candidates:
                        seen_candidates.add(p_val)
                        candidates.append(p_val)
            elif len(cells) >= 2:
                c0, c1 = cells[0].strip(), cells[1].strip()
                if c0 and c1:
                    c_list, m_dict = extract_passwords_from_text(f"{c0}: {c1}")
                    merge_results(c_list, m_dict)
            elif len(cells) == 1:
                c_list, m_dict = extract_passwords_from_text(cells[0])
                merge_results(c_list, m_dict)

    # 2. Inspect paragraphs
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            c_list, m_dict = extract_passwords_from_text(text)
            merge_results(c_list, m_dict)

    return candidates, name_map


def is_credential_docx(filename: str, file_bytes: bytes) -> bool:
    """
    Determines if a .docx document is a password/credential file rather than an audit report.
    Checks filename keywords and content brevity / presence of password mappings.
    """
    base_name = os.path.basename(filename).lower()
    for kw in PWD_FILE_KEYWORDS:
        if kw in base_name:
            return True

    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        full_text = " ".join(p.text for p in doc.paragraphs).lower()
        # If very short (< 400 chars) and mentions password/pwd
        if len(full_text) < 400 and any(kw in full_text for kw in ("password", "pwd", "passcode", "secret")):
            # If it lacks audit reporting markers, treat as credential file
            if not any(term in full_text for term in ("vulnerability", "cve", "cvss", "observation", "audit report", "executive summary", "scope")):
                return True
    except Exception:
        pass

    return False


def decode_text_bytes(raw_bytes: bytes) -> str:
    """Safely decodes text bytes across various encodings."""
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw_bytes.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw_bytes.decode("utf-8", errors="ignore")


def analyze_and_extract_zip(
    zip_bytes: bytes,
    user_password_map: Optional[dict] = None,
    user_bulk_passwords: Optional[list] = None,
) -> Tuple[List[ZipExtractedReport], Dict[str, str], List[str]]:
    """
    Unpacks a ZIP archive in memory, extracts passwords from text and docx files,
    pre-emptively unlocks protected reports, and returns valid reports ready for batch processing.

    Returns:
        (extracted_reports, global_name_password_map, global_candidate_passwords)
    """
    zip_io = io.BytesIO(zip_bytes)
    if not zipfile.is_zipfile(zip_io):
        raise ValueError("Uploaded file is not a valid ZIP archive.")

    global_name_map: Dict[str, str] = {}
    if user_password_map:
        for k, v in user_password_map.items():
            if k and v:
                global_name_map[str(k).strip().lower()] = str(v).strip()

    global_passwords: List[str] = []
    seen_global_pwds: Set[str] = set()

    def add_global_password(p: str):
        p = p.strip()
        if p and p not in seen_global_pwds:
            seen_global_pwds.add(p)
            global_passwords.append(p)

    if user_bulk_passwords:
        for p in user_bulk_passwords:
            if p:
                add_global_password(str(p))

    # Directory-scoped passwords: normalized_dir_path -> list of candidate passwords
    dir_passwords: Dict[str, List[str]] = {}

    def add_dir_password(dir_p: str, p: str):
        p = p.strip()
        if not p:
            return
        if dir_p not in dir_passwords:
            dir_passwords[dir_p] = []
        if p not in dir_passwords[dir_p]:
            dir_passwords[dir_p].append(p)
        add_global_password(p)

    report_candidates: List[Tuple[str, str, str, bytes]] = []  # (rel_path, dir_path, base_name, bytes)

    with zipfile.ZipFile(zip_io, "r") as zf:
        # Pass 1: Scan all entries for password files & report candidates
        for member in zf.infolist():
            if member.is_dir():
                continue

            # Zip-slip prevention & path normalization
            raw_name = member.filename.replace("\\", "/").lstrip("/")
            path_parts = [p for p in raw_name.split("/") if p and p != "."]
            if ".." in path_parts:
                logger.warning(f"Skipping dangerous path traversal entry in zip: {member.filename}")
                continue

            rel_path = "/".join(path_parts)
            base_name = os.path.basename(rel_path)
            dir_path = os.path.dirname(rel_path)

            # Skip OS metadata and hidden files
            if base_name.startswith(".") or base_name.startswith("~") or rel_path.startswith("__MACOSX/"):
                continue

            ext = os.path.splitext(base_name)[1].lower()
            file_bytes = zf.read(member)

            # 1. Text files (.txt, .csv, .tsv)
            if ext in VALID_TEXT_EXTENSIONS:
                text_content = decode_text_bytes(file_bytes)
                cands, n_map = extract_passwords_from_text(text_content)
                for c in cands:
                    add_dir_password(dir_path, c)
                for k, v in n_map.items():
                    global_name_map[k] = v
                    add_dir_password(dir_path, v)
                logger.info(f"Harvested {len(cands)} passwords from text file: {rel_path}")

            # 2. Word documents (.docx)
            elif ext == ".docx":
                if is_credential_docx(base_name, file_bytes):
                    cands, n_map = extract_passwords_from_docx(file_bytes)
                    for c in cands:
                        add_dir_password(dir_path, c)
                    for k, v in n_map.items():
                        global_name_map[k] = v
                        add_dir_password(dir_path, v)
                    logger.info(f"Harvested {len(cands)} passwords from credential Word file: {rel_path}")
                else:
                    # Valid report candidate
                    report_candidates.append((rel_path, dir_path, base_name, file_bytes))

            # 3. PDF and DOC files
            elif ext in {".pdf", ".doc"}:
                report_candidates.append((rel_path, dir_path, base_name, file_bytes))

        # Pass 2: Decrypt & prepare report candidates
        final_reports: List[ZipExtractedReport] = []

        for rel_path, dir_path, base_name, file_bytes in report_candidates:
            ext = os.path.splitext(base_name)[1].lower()
            stem_name = os.path.splitext(base_name)[0].lower()

            is_enc = False
            if ext == ".pdf":
                is_enc = is_pdf_encrypted(file_bytes)
            elif ext == ".docx":
                is_enc = is_docx_encrypted(file_bytes)

            if not is_enc:
                final_reports.append(
                    ZipExtractedReport(
                        relative_path=rel_path,
                        filename=base_name,
                        file_bytes=file_bytes,
                        protection_status="No",
                    )
                )
                continue

            # Prioritize candidate passwords for this specific report
            candidates_to_try: List[str] = []
            seen_to_try: Set[str] = set()

            def queue_pwd(pwd: Optional[str]):
                if pwd and pwd not in seen_to_try:
                    seen_to_try.add(pwd)
                    candidates_to_try.append(pwd)

            # Priority 1: Exact filename match (e.g. audit_q1.pdf or audit_q1)
            queue_pwd(global_name_map.get(stem_name))
            queue_pwd(global_name_map.get(base_name.lower()))
            queue_pwd(global_name_map.get(rel_path.lower()))

            # Priority 2: Same directory passwords
            for p in dir_passwords.get(dir_path, []):
                queue_pwd(p)

            # Priority 3: Ancestor directory passwords
            curr_dir = dir_path
            while curr_dir:
                parent_dir = os.path.dirname(curr_dir)
                for p in dir_passwords.get(parent_dir, []):
                    queue_pwd(p)
                if not parent_dir or parent_dir == curr_dir:
                    break
                curr_dir = parent_dir

            # Priority 4: Root directory / global archive passwords
            for p in dir_passwords.get("", []):
                queue_pwd(p)
            for p in global_passwords:
                queue_pwd(p)

            # Attempt decryption with ordered candidates
            decrypted = False
            decrypted_bytes = file_bytes

            for candidate in candidates_to_try:
                try:
                    if ext == ".pdf":
                        decrypted_bytes = decrypt_pdf(file_bytes, candidate)
                    elif ext == ".docx":
                        decrypted_bytes = decrypt_docx(file_bytes, candidate)
                    decrypted = True
                    logger.info(f"Successfully decrypted {rel_path} using harvested password")
                    break
                except Exception:
                    continue

            if decrypted:
                final_reports.append(
                    ZipExtractedReport(
                        relative_path=rel_path,
                        filename=base_name,
                        file_bytes=decrypted_bytes,
                        protection_status="Decrypted",
                    )
                )
            else:
                logger.warning(f"Could not decrypt protected file {rel_path} with available passwords")
                final_reports.append(
                    ZipExtractedReport(
                        relative_path=rel_path,
                        filename=base_name,
                        file_bytes=file_bytes,
                        protection_status="Locked",
                    )
                )

    return final_reports, global_name_map, global_passwords
